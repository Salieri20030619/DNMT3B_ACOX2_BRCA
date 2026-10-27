#!/usr/bin/env python3
from __future__ import annotations

import csv
import gzip
import os
import re
import sys
import zipfile
from collections import Counter, defaultdict
from pathlib import Path


PROJECT_ROOT = Path(
    os.environ.get("PROJECT_ROOT", "/mnt/BioDisk/wangmingyuan/DNMT3B_ACOX2_BRCA")
)
RAW_DIR = PROJECT_ROOT / "00_raw_data"
RUN_DIR = Path(os.environ["RUN_DIR"])

FILES = {
    "expr": RAW_DIR / "TCGA_BRCA/TCGA.BRCA.HiSeqV2.tsv.gz",
    "clinical": RAW_DIR / "TCGA_BRCA/TCGA.BRCA.clinical.tsv",
    "methylation": RAW_DIR / "TCGA_BRCA/TCGA.BRCA.HumanMethylation450.tsv.gz",
    "cdr": RAW_DIR / "TCGA_BRCA/TCGA-CDR-SupplementalTableS1.xlsx",
    "gse70947": RAW_DIR / "GEO/GSE70947/GSE70947_series_matrix.txt.gz",
    "gse65194": RAW_DIR / "GEO/GSE65194/GSE65194_series_matrix.txt.gz",
    "gpl13607": RAW_DIR / "GEO/platforms/GPL13607_old_annotations.txt.gz",
    "gpl570": RAW_DIR / "GEO/platforms/GPL570.annot.gz",
}

MISSING_VALUES = {
    "",
    "na",
    "n/a",
    "null",
    "unknown",
    "not reported",
    "[not available]",
    "[not evaluated]",
    "[unknown]",
    "--",
}


def write_tsv(name: str, rows: list[dict], fieldnames: list[str] | None = None) -> None:
    path = RUN_DIR / name
    if fieldnames is None:
        fieldnames = list(rows[0].keys()) if rows else ["status"]
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames, delimiter="\t", extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)


def strip_quotes(value: str) -> str:
    value = value.strip()
    if len(value) >= 2 and value[0] == '"' and value[-1] == '"':
        value = value[1:-1]
    return value.replace('""', '"')


def normalize_sample(value: str) -> str:
    value = strip_quotes(value).upper()
    match = re.match(r"^(TCGA-[A-Z0-9]{2}-[A-Z0-9]{4}-[0-9]{2})", value)
    return match.group(1) if match else value


def patient_id(value: str) -> str | None:
    value = normalize_sample(value)
    return value[:12] if value.startswith("TCGA-") and len(value) >= 12 else None


def sample_code(value: str) -> str | None:
    value = normalize_sample(value)
    return value[13:15] if value.startswith("TCGA-") and len(value) >= 15 else None


def valid_value(value: str | None) -> bool:
    if value is None:
        return False
    return str(value).strip().lower() not in MISSING_VALUES


def count_paired_patients(ids: list[str]) -> int:
    codes_by_patient: dict[str, set[str]] = defaultdict(set)
    for sample in ids:
        patient = patient_id(sample)
        code = sample_code(sample)
        if patient and code:
            codes_by_patient[patient].add(code)
    return sum({"01", "11"}.issubset(codes) for codes in codes_by_patient.values())


def audit_gzip_matrix(
    path: Path, target_features: tuple[str, ...] = ()
) -> tuple[list[str], int, dict[str, dict[str, int]]]:
    targets = {feature.upper() for feature in target_features}
    target_stats: dict[str, dict[str, int]] = {}
    with gzip.open(path, "rt", encoding="utf-8", errors="replace", newline="") as handle:
        header_line = handle.readline()
        if not header_line:
            raise RuntimeError(f"Empty gzip matrix: {path}")
        header = [strip_quotes(x) for x in header_line.rstrip("\r\n").split("\t")]
        row_count = 0
        for line in handle:
            row_count += 1
            if not targets:
                continue
            fields = line.rstrip("\r\n").split("\t")
            feature = strip_quotes(fields[0]).upper()
            if feature in targets:
                values = [strip_quotes(value) for value in fields[1:]]
                target_stats[feature] = {
                    "values": len(values),
                    "valid_values": sum(valid_value(value) for value in values),
                }
    return header, row_count, target_stats


def modality_row(modality: str, ids: list[str], feature_count: int) -> dict:
    normalized = [normalize_sample(x) for x in ids]
    codes = [sample_code(x) for x in normalized]
    patients = {x for x in (patient_id(y) for y in normalized) if x}
    return {
        "modality": modality,
        "features": feature_count,
        "samples": len(normalized),
        "unique_samples": len(set(normalized)),
        "unique_patients": len(patients),
        "primary_tumor": sum(x == "01" for x in codes),
        "metastatic": sum(x == "06" for x in codes),
        "solid_tissue_normal": sum(x == "11" for x in codes),
        "paired_patients_01_11": count_paired_patients(normalized),
    }


def read_clinical(path: Path) -> tuple[list[dict], list[str]]:
    with path.open("r", encoding="utf-8", errors="replace", newline="") as handle:
        reader = csv.DictReader(handle, delimiter="\t")
        rows = list(reader)
        fields = reader.fieldnames or []
    if "sampleID" not in fields:
        raise RuntimeError("sampleID is absent from TCGA clinical matrix")
    for row in rows:
        row["sample15"] = normalize_sample(row["sampleID"])
        row["patient12"] = patient_id(row["sampleID"]) or ""
        row["sample_code"] = sample_code(row["sampleID"]) or ""
    return rows, fields


def distribution(rows: list[dict], column: str) -> list[dict]:
    counts = Counter(row.get(column, "") for row in rows if valid_value(row.get(column)))
    return [{column: key, "n": value} for key, value in counts.most_common()]


def read_platform_targets(path: Path, platform: str) -> list[dict]:
    hits: list[dict] = []
    with gzip.open(path, "rt", encoding="utf-8", errors="replace") as handle:
        for line in handle:
            if line.startswith(("!", "#")) or "\t" not in line:
                continue
            upper = line.upper()
            genes = [gene for gene in ("ACOX2", "DNMT3B") if gene in upper]
            if not genes:
                continue
            fields = line.rstrip("\r\n").split("\t")
            probe = fields[0] if fields else ""
            for gene in genes:
                hits.append(
                    {
                        "platform": platform,
                        "gene": gene,
                        "probe_id": probe,
                        "raw_record": line.rstrip("\r\n"),
                    }
                )
    unique: dict[tuple[str, str, str], dict] = {}
    for row in hits:
        unique[(row["platform"], row["gene"], row["probe_id"])] = row
    return list(unique.values())


def derive_pair_key(title: str) -> str:
    value = title.lower()
    value = re.sub(r"normal|tumou?r|adenocarcinoma|adjacent|breast|tissue|sample", "", value)
    value = re.sub(r"[^a-z0-9]", "", value)
    value = re.sub(r"^[nt]([0-9]+)$", r"\1", value)
    value = re.sub(r"^([0-9]+)[nt]$", r"\1", value)
    return value


def resolve_field(fields: list[str], candidates: tuple[str, ...]) -> str | None:
    exact = {field.lower(): field for field in fields}
    for candidate in candidates:
        if candidate.lower() in exact:
            return exact[candidate.lower()]
    for candidate in candidates:
        pattern = re.compile(candidate, re.IGNORECASE)
        matches = [field for field in fields if pattern.search(field)]
        if matches:
            return matches[0]
    return None


def classify_geo(dataset: str, text: str) -> str:
    text = text.lower()
    normal = re.search(
        r"normal breast|adjacent normal|non[- _]?tumou?r|healthy breast|normal tissue", text
    )
    if normal:
        return "Normal"
    if dataset == "GSE65194":
        if re.search(r"non[- _]?tnbc|non[- _]?triple", text):
            return "Non_TNBC"
        if re.search(r"tnbc|triple[- _]?negative", text):
            return "TNBC"
    if re.search(r"tumou?r|cancer|adenocarcinoma|carcinoma", text):
        return "Tumor" if dataset == "GSE70947" else "Tumor_other"
    return "Unresolved"


def scan_geo_matrix(path: Path, dataset: str, target_probes: set[str]) -> dict:
    metadata_rows: list[tuple[str, list[str]]] = []
    in_table = False
    table_header_seen = False
    feature_count = 0
    present_probes: set[str] = set()

    with gzip.open(path, "rt", encoding="utf-8", errors="replace") as handle:
        for raw_line in handle:
            line = raw_line.rstrip("\r\n")
            if line == "!series_matrix_table_begin":
                in_table = True
                continue
            if line == "!series_matrix_table_end":
                in_table = False
                continue
            if not in_table and line.startswith("!Sample_"):
                fields = line.split("\t")
                metadata_rows.append((fields[0], [strip_quotes(x) for x in fields[1:]]))
                continue
            if in_table:
                if not table_header_seen:
                    table_header_seen = True
                    continue
                probe = strip_quotes(line.split("\t", 1)[0])
                feature_count += 1
                if probe in target_probes:
                    present_probes.add(probe)

    accession_rows = [values for tag, values in metadata_rows if tag == "!Sample_geo_accession"]
    if len(accession_rows) != 1:
        raise RuntimeError(f"Unexpected !Sample_geo_accession rows in {dataset}: {len(accession_rows)}")
    accessions = accession_rows[0]
    sample_count = len(accessions)
    titles = [""] * sample_count
    combined = [""] * sample_count

    for tag, values in metadata_rows:
        if len(values) != sample_count:
            continue
        if tag == "!Sample_title":
            titles = values
        if tag in {"!Sample_title", "!Sample_source_name_ch1", "!Sample_characteristics_ch1"}:
            combined = [f"{old} | {new}" for old, new in zip(combined, values)]

    rows = []
    for accession, title, text in zip(accessions, titles, combined):
        rows.append(
            {
                "dataset": dataset,
                "sample": accession,
                "title": title,
                "provisional_group": classify_geo(dataset, text),
                "provisional_pair_key": derive_pair_key(title),
                "combined_metadata": text.strip(" |"),
            }
        )
    return {
        "metadata": rows,
        "features": feature_count,
        "present_probes": present_probes,
    }


def audit_cdr(path: Path) -> list[dict]:
    result = [
        {"metric": "xlsx_zip_valid", "value": zipfile.is_zipfile(path)},
        {"metric": "openpyxl_available", "value": False},
        {"metric": "workbook_read", "value": False},
        {"metric": "BRCA_rows", "value": "NA"},
    ]
    try:
        import openpyxl  # type: ignore
    except ImportError:
        return result
    result[1]["value"] = True
    try:
        workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
        sheet = workbook[workbook.sheetnames[0]]
        iterator = sheet.iter_rows(values_only=True)
        header = [str(x) if x is not None else "" for x in next(iterator)]
        candidate = next(
            (name for name in ("type", "acronym", "project_id", "cancer_type") if name in header),
            None,
        )
        brca_rows = 0
        if candidate:
            index = header.index(candidate)
            for row in iterator:
                if index < len(row) and str(row[index]).upper() == "BRCA":
                    brca_rows += 1
        result[2]["value"] = True
        result[3]["value"] = brca_rows if candidate else "NA"
        workbook.close()
    except Exception as exc:  # noqa: BLE001
        print(f"CDR_READ_WARNING={exc}")
    return result


def main() -> int:
    RUN_DIR.mkdir(parents=True, exist_ok=True)
    for label, path in FILES.items():
        if not path.is_file() or path.stat().st_size == 0:
            raise RuntimeError(f"Required file missing or empty: {label}={path}")

    print("PY_STEP=TCGA_MATRIX_AUDIT", flush=True)
    expr_header, expr_features, target_expression = audit_gzip_matrix(
        FILES["expr"], ("DNMT3B", "ACOX2")
    )
    meth_header, meth_probes, _ = audit_gzip_matrix(FILES["methylation"])
    missing_targets = sorted({"DNMT3B", "ACOX2"} - set(target_expression))
    if missing_targets:
        raise RuntimeError(f"Target expression rows absent: {','.join(missing_targets)}")
    expr_ids = expr_header[1:]
    meth_ids = meth_header[1:]
    modality_rows = [
        modality_row("RNA_HiSeqV2", expr_ids, expr_features),
        modality_row("Methylation450", meth_ids, meth_probes),
    ]
    write_tsv("tcga_modality_counts.tsv", modality_rows)
    write_tsv(
        "tcga_target_expression_rows.tsv",
        [
            {
                "gene": gene,
                "row_found": gene in target_expression,
                "values": target_expression.get(gene, {}).get("values", 0),
                "valid_values": target_expression.get(gene, {}).get("valid_values", 0),
            }
            for gene in ("DNMT3B", "ACOX2")
        ],
    )

    print("PY_STEP=TCGA_CLINICAL_AUDIT", flush=True)
    clinical, clinical_fields = read_clinical(FILES["clinical"])
    primary_clinical = [row for row in clinical if row["sample_code"] == "01"]
    requested_fields = {
        "sample_type": ("sample_type", r"sample.*type"),
        "PAM50": ("PAM50_mRNA_nature2012", "PAM50Call_RNAseq", r"PAM50"),
        "AJCC_stage": ("AJCC_Stage_nature2012", r"AJCC.*stage", r"pathologic.*stage"),
        "ER_status": ("ER_Status_nature2012", r"(^|_)ER.*status"),
        "PR_status": ("PR_Status_nature2012", r"(^|_)PR.*status"),
        "HER2_status": ("HER2_Final_Status_nature2012", r"HER2.*status"),
    }
    field_rows = []
    resolved_fields = {
        label: resolve_field(clinical_fields, candidates)
        for label, candidates in requested_fields.items()
    }
    for label, candidates in requested_fields.items():
        column = resolved_fields[label]
        found = column is not None
        field_rows.append(
            {
                "field": label,
                "requested_column": candidates[0],
                "resolved_column": column or "NA",
                "found": found,
                "valid_all": sum(valid_value(row.get(column)) for row in clinical) if found else 0,
                "valid_primary_tumor": (
                    sum(valid_value(row.get(column)) for row in primary_clinical) if found else 0
                ),
            }
        )
    write_tsv("tcga_clinical_field_audit.tsv", field_rows)
    write_tsv(
        "tcga_clinical_counts.tsv",
        [
            {"metric": "clinical_rows", "value": len(clinical)},
            {"metric": "clinical_unique_samples", "value": len({x["sample15"] for x in clinical})},
            {
                "metric": "clinical_unique_patients",
                "value": len({x["patient12"] for x in clinical if x["patient12"]}),
            },
            {"metric": "clinical_primary_tumor_rows", "value": len(primary_clinical)},
        ],
    )
    pam50_column = resolved_fields["PAM50"]
    stage_column = resolved_fields["AJCC_stage"]
    if pam50_column:
        write_tsv(
            "tcga_PAM50_distribution.tsv",
            distribution(primary_clinical, pam50_column),
            [pam50_column, "n"],
        )
    if stage_column:
        write_tsv(
            "tcga_AJCC_stage_distribution.tsv",
            distribution(primary_clinical, stage_column),
            [stage_column, "n"],
        )

    print("PY_STEP=TCGA_INTERSECTIONS", flush=True)
    expr15 = {normalize_sample(x) for x in expr_ids}
    meth15 = {normalize_sample(x) for x in meth_ids}
    clin15 = {row["sample15"] for row in clinical}
    expr_patients = {x for x in (patient_id(y) for y in expr15) if x}
    meth_patients = {x for x in (patient_id(y) for y in meth15) if x}
    clin_patients = {row["patient12"] for row in clinical if row["patient12"]}
    primary_overlap = {x for x in expr15 & meth15 if sample_code(x) == "01"}
    normal_overlap = {x for x in expr15 & meth15 if sample_code(x) == "11"}
    intersection_rows = [
        {"metric": "RNA_Methylation_exact_samples", "value": len(expr15 & meth15)},
        {"metric": "RNA_Clinical_exact_samples", "value": len(expr15 & clin15)},
        {"metric": "Methylation_Clinical_exact_samples", "value": len(meth15 & clin15)},
        {"metric": "RNA_Methylation_Clinical_exact_samples", "value": len(expr15 & meth15 & clin15)},
        {"metric": "RNA_Methylation_patients", "value": len(expr_patients & meth_patients)},
        {
            "metric": "RNA_Methylation_Clinical_patients",
            "value": len(expr_patients & meth_patients & clin_patients),
        },
        {"metric": "RNA_Methylation_primary_tumor_samples", "value": len(primary_overlap)},
        {"metric": "RNA_Methylation_normal_samples", "value": len(normal_overlap)},
    ]
    write_tsv("tcga_intersection_audit.tsv", intersection_rows)

    print("PY_STEP=TCGA_CDR_AUDIT", flush=True)
    write_tsv("tcga_CDR_audit.tsv", audit_cdr(FILES["cdr"]))

    print("PY_STEP=GEO_PLATFORM_TARGETS", flush=True)
    platform_targets = read_platform_targets(FILES["gpl13607"], "GPL13607")
    platform_targets += read_platform_targets(FILES["gpl570"], "GPL570")
    write_tsv("geo_target_probe_annotation.tsv", platform_targets)
    probes13607 = {x["probe_id"] for x in platform_targets if x["platform"] == "GPL13607"}
    probes570 = {x["probe_id"] for x in platform_targets if x["platform"] == "GPL570"}

    print("PY_STEP=GEO_MATRIX_AUDIT", flush=True)
    geo70947 = scan_geo_matrix(FILES["gse70947"], "GSE70947", probes13607)
    geo65194 = scan_geo_matrix(FILES["gse65194"], "GSE65194", probes570)
    geo_metadata = geo70947["metadata"] + geo65194["metadata"]
    write_tsv("geo_sample_metadata_audit.tsv", geo_metadata)
    group_counts = Counter((x["dataset"], x["provisional_group"]) for x in geo_metadata)
    write_tsv(
        "geo_group_counts.tsv",
        [
            {"dataset": dataset, "provisional_group": group, "n": count}
            for (dataset, group), count in sorted(group_counts.items())
        ],
    )
    write_tsv(
        "geo_matrix_audit.tsv",
        [
            {
                "dataset": "GSE70947",
                "samples": len(geo70947["metadata"]),
                "features": geo70947["features"],
                "target_probes_expected": len(probes13607),
                "target_probes_present": len(geo70947["present_probes"]),
            },
            {
                "dataset": "GSE65194",
                "samples": len(geo65194["metadata"]),
                "features": geo65194["features"],
                "target_probes_expected": len(probes570),
                "target_probes_present": len(geo65194["present_probes"]),
            },
        ],
    )
    probe_gene = {(x["platform"], x["probe_id"]): x["gene"] for x in platform_targets}
    probe_presence = []
    for dataset, platform, expected, present in (
        ("GSE70947", "GPL13607", probes13607, geo70947["present_probes"]),
        ("GSE65194", "GPL570", probes570, geo65194["present_probes"]),
    ):
        for probe in sorted(expected):
            probe_presence.append(
                {
                    "dataset": dataset,
                    "platform": platform,
                    "gene": probe_gene.get((platform, probe), "NA"),
                    "probe_id": probe,
                    "present": probe in present,
                }
            )
    write_tsv("geo_target_probe_presence.tsv", probe_presence)

    g709 = geo70947["metadata"]
    groups_by_key: dict[str, set[str]] = defaultdict(set)
    for row in g709:
        if row["provisional_pair_key"]:
            groups_by_key[row["provisional_pair_key"]].add(row["provisional_group"])
    provisional_pairs = sum({"Tumor", "Normal"}.issubset(groups) for groups in groups_by_key.values())

    promoter_status = "PENDING_450K_ANNOTATION_PACKAGE_CHECK"
    write_tsv(
        "acox2_promoter_probe_audit.tsv",
        [
            {"metric": "target_gene", "value": "ACOX2"},
            {"metric": "target_transcript", "value": "NM_003500.4"},
            {"metric": "methylation_probe_rows", "value": meth_probes},
            {"metric": "promoter_probe_mapping", "value": promoter_status},
        ],
    )

    summary = [
        "AUDIT_SCOPE=TCGA_BRCA_GEO_ACOX2",
        f"TCGA_RNA_FEATURES={expr_features}",
        f"TCGA_RNA_SAMPLES={len(expr_ids)}",
        f"TCGA_RNA_PRIMARY_TUMOR={sum(sample_code(x) == '01' for x in expr_ids)}",
        f"TCGA_RNA_NORMAL={sum(sample_code(x) == '11' for x in expr_ids)}",
        f"TCGA_RNA_PAIRED_PATIENTS={count_paired_patients(expr_ids)}",
        f"TCGA_METHYLATION_PROBES={meth_probes}",
        f"TCGA_METHYLATION_SAMPLES={len(meth_ids)}",
        f"TCGA_RNA_METHYLATION_EXACT_OVERLAP={len(expr15 & meth15)}",
        f"TCGA_PRIMARY_RNA_METHYLATION_OVERLAP={len(primary_overlap)}",
        f"GSE70947_SAMPLES={len(geo70947['metadata'])}",
        f"GSE70947_FEATURES={geo70947['features']}",
        f"GSE70947_PROVISIONAL_PAIRED_KEYS={provisional_pairs}",
        f"GSE65194_SAMPLES={len(geo65194['metadata'])}",
        f"GSE65194_FEATURES={geo65194['features']}",
        f"ACOX2_PROMOTER_AUDIT_STATUS={promoter_status}",
        "RAW_DATA_MODIFIED=NO",
        "AUDIT_PASS",
    ]
    (RUN_DIR / "audit_summary.txt").write_text("\n".join(summary) + "\n", encoding="utf-8")
    (RUN_DIR / "python_runtime.txt").write_text(sys.version + "\n", encoding="utf-8")
    print("\n".join(summary))
    return 0


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except Exception as error:  # noqa: BLE001
        print(f"AUDIT_ERROR={type(error).__name__}:{error}", file=sys.stderr)
        raise
